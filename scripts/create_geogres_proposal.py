from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.shared import RGBColor
from docx.oxml.ns import qn

from create_autozentic_proposal import (
    BLACK,
    BODY,
    BORDER,
    LIGHT_GRAY,
    LIGHT_TEAL,
    NAVY,
    WHITE,
    add_body,
    add_bullet,
    add_heading,
    add_label_value_row,
    add_page_number,
    add_run,
    configure_document,
    set_cell_shading,
    set_cell_text,
    set_repeat_table_header,
    set_table_borders,
    set_cell_margins,
)


OUT = Path(__file__).resolve().parent.parent / "Autozentic_Geogres_CRM_Proposal.docx"


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(4)
    add_run(title, "CRM and Social Chatbot Proposal", bold=True, color=BLACK, size=27)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    add_run(subtitle, "24/7 customer replies, lead management and follow-up automation", color=RGBColor(83, 96, 112), size=11.5)

    meta = doc.add_table(rows=0, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    meta.columns[0].width = Inches(1.35)
    meta.columns[1].width = Inches(5.15)
    set_table_borders(meta, color="FFFFFF", size="0")
    add_label_value_row(meta, "Prepared by", "Autozentic")
    add_label_value_row(meta, "Location", "Jagatpura, Jaipur")
    add_label_value_row(meta, "Prepared for", "Geogres")
    add_label_value_row(meta, "Date", "8 September 2026")
    for row in meta.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=45, bottom=45, start=0, end=120)

    add_heading(doc, "Proposal Summary", before=22, after=6)
    add_body(doc, "Autozentic will configure a customer-engagement CRM for Geogres with 24/7 conversational chatbots across WhatsApp, Instagram and Facebook. The solution will organize incoming enquiries, move leads through a clear pipeline, manage dealers and partners, schedule follow-ups and keep appointments visible to the team.", after=7)

    add_heading(doc, "Included Solution", before=16, after=7)
    scope = doc.add_table(rows=1, cols=2)
    scope.alignment = WD_TABLE_ALIGNMENT.LEFT
    scope.autofit = False
    scope.columns[0].width = Inches(1.95)
    scope.columns[1].width = Inches(4.55)
    set_table_borders(scope)
    for cell in scope.rows[0].cells:
        set_cell_shading(cell, NAVY)
    set_cell_text(scope.rows[0].cells[0], "Module", bold=True, color=WHITE, size=9.2)
    set_cell_text(scope.rows[0].cells[1], "Included capability", bold=True, color=WHITE, size=9.2)
    set_repeat_table_header(scope.rows[0])
    rows = [
        ("24/7 social chatbots", "Connected WhatsApp, Instagram and Facebook chat experiences for FAQs, enquiry handling, lead capture and customer responses at any time."),
        ("CRM dashboard", "A central workspace for activity visibility, lead status, appointments, follow-ups and operational summaries."),
        ("Leads pipeline", "Capture, qualify, assign and move leads through configurable stages with customer history available to the team."),
        ("Dealers and partners", "Dealer and partner records, contact details, relationship tracking and follow-up visibility."),
        ("Follow-ups", "Scheduled follow-up tasks and reminders so enquiries and opportunities do not get missed."),
        ("Appointments", "Appointment creation, date and time tracking, status updates and visibility for the responsible team member."),
    ]
    for index, (label, description) in enumerate(rows):
        row = scope.add_row()
        if index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
        set_cell_text(row.cells[0], label, bold=True, color=BLACK, size=9.1)
        set_cell_text(row.cells[1], description, size=9.1)

    add_heading(doc, "Delivery Approach", before=16, after=5)
    add_bullet(doc, "Confirm Geogres workflows, user roles, lead stages, appointment rules and response content.")
    add_bullet(doc, "Connect and configure WhatsApp, Instagram and Facebook chatbot entry points.")
    add_bullet(doc, "Configure the CRM dashboard, lead pipeline, dealer and partner records, follow-ups and appointments.")
    add_bullet(doc, "Test the workflows with sample enquiries and complete handover within the agreed seven-day timeline.")

    doc.add_page_break()
    add_heading(doc, "Timeline", before=0, after=7)
    add_body(doc, "The target delivery timeline is seven calendar days from receipt of the advance payment and the required account access, approvals and business information.", after=10)

    timeline = doc.add_table(rows=1, cols=2)
    timeline.alignment = WD_TABLE_ALIGNMENT.LEFT
    timeline.autofit = False
    timeline.columns[0].width = Inches(1.55)
    timeline.columns[1].width = Inches(4.95)
    set_table_borders(timeline)
    for cell, text in zip(timeline.rows[0].cells, ("Stage", "Planned work")):
        set_cell_shading(cell, NAVY)
        set_cell_text(cell, text, bold=True, color=WHITE, size=9.2)
    set_repeat_table_header(timeline.rows[0])
    stages = [
        ("Days 1–2", "Access, workflow confirmation, chatbot flow and CRM structure."),
        ("Days 3–5", "Channel connections, dashboard, pipeline, dealer and partner setup, follow-up and appointment configuration."),
        ("Days 6–7", "Testing, corrections, client review, handover and usage guidance."),
    ]
    for index, values in enumerate(stages):
        row = timeline.add_row()
        if index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
        set_cell_text(row.cells[0], values[0], bold=True, color=BLACK, size=9.2)
        set_cell_text(row.cells[1], values[1], size=9.2)

    add_heading(doc, "Commercials", before=18, after=7)
    add_body(doc, "The complete project is offered at a total professional fee of INR 35,000. This proposal covers the configuration and delivery described above.", after=10)

    pricing = doc.add_table(rows=1, cols=2)
    pricing.alignment = WD_TABLE_ALIGNMENT.LEFT
    pricing.autofit = False
    pricing.columns[0].width = Inches(4.75)
    pricing.columns[1].width = Inches(1.75)
    set_table_borders(pricing)
    for cell, text in zip(pricing.rows[0].cells, ("Commercial item", "Amount")):
        set_cell_shading(cell, NAVY)
        set_cell_text(cell, text, bold=True, color=WHITE, size=9.2)
    set_repeat_table_header(pricing.rows[0])
    for index, values in enumerate((("Complete CRM and chatbot solution", "INR 35,000"), ("Total project fee", "INR 35,000"))):
        row = pricing.add_row()
        if index == 1:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_TEAL)
        set_cell_text(row.cells[0], values[0], bold=index == 1, color=BLACK if index == 1 else BODY, size=9.4)
        set_cell_text(row.cells[1], values[1], bold=True, color=BLACK, size=9.4, align=WD_ALIGN_PARAGRAPH.RIGHT)

    add_heading(doc, "Payment Schedule", before=16, after=6)
    payment = doc.add_table(rows=1, cols=3)
    payment.alignment = WD_TABLE_ALIGNMENT.LEFT
    payment.autofit = False
    payment.columns[0].width = Inches(1.65)
    payment.columns[1].width = Inches(3.15)
    payment.columns[2].width = Inches(1.7)
    set_table_borders(payment)
    for cell, text in zip(payment.rows[0].cells, ("Milestone", "Payment condition", "Amount")):
        set_cell_shading(cell, NAVY)
        set_cell_text(cell, text, bold=True, color=WHITE, size=9.1)
    set_repeat_table_header(payment.rows[0])
    for index, values in enumerate((("Advance", "50% before project commencement", "INR 17,500"), ("Completion", "50% after completion and handover", "INR 17,500"))):
        row = payment.add_row()
        if index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
        set_cell_text(row.cells[0], values[0], bold=True, color=BLACK, size=9.2)
        set_cell_text(row.cells[1], values[1], size=9.2)
        set_cell_text(row.cells[2], values[2], size=9.2, align=WD_ALIGN_PARAGRAPH.RIGHT)

    add_heading(doc, "Client Paid Provider Charges", before=16, after=6)
    add_body(doc, "Charges billed by WhatsApp, Meta, messaging providers or other third-party services are separate from Autozentic's professional fee and are paid directly by the client. Provider account approvals and usage charges remain the client's responsibility.", after=7)

    doc.add_page_break()
    add_heading(doc, "Support and Responsibilities", before=0, after=7)
    add_heading(doc, "Included support", level=2, before=4, after=5)
    add_body(doc, "Autozentic will provide one year of bug and technical support for the delivered CRM and chatbot implementation. Bug support covers defects in the delivered configuration and integrations. New features, new workflows, new content and provider charges may be quoted separately.", after=9)

    add_heading(doc, "Client responsibilities", level=2, before=12, after=5)
    add_bullet(doc, "Provide access to the required WhatsApp, Instagram, Facebook, CRM and business accounts.")
    add_bullet(doc, "Complete provider verification, approvals and payment of third-party usage charges.")
    add_bullet(doc, "Provide approved business information, FAQs, lead stages, appointment rules and escalation contacts.")
    add_bullet(doc, "Review and approve the configured workflows during testing and handover.")

    closing = add_body(doc, "We are looking forward to working further with you.", after=0)
    closing.paragraph_format.space_before = Pt(18)

    doc.core_properties.title = "CRM and Social Chatbot Proposal - Geogres"
    doc.core_properties.subject = "Proposal for 24/7 social chatbots and CRM workflow automation"
    doc.core_properties.author = "Autozentic"
    doc.core_properties.keywords = "Autozentic, Geogres, CRM, WhatsApp, Instagram, Facebook, chatbot, proposal"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
