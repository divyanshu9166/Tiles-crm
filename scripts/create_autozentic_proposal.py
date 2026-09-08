from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent.parent / "Autozentic_AI_Automation_Proposal.docx"

NAVY = "18324B"
TEAL = "0F766E"
LIGHT_TEAL = "EAF5F3"
LIGHT_GRAY = "F4F6F8"
BORDER = "D9DEE5"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
BODY = RGBColor(42, 50, 61)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, color=BODY, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_run(paragraph, text, *, bold=False, italic=False, color=BODY, size=10.5):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return run


def add_body(doc, text, *, after=7, size=10.5):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    add_run(p, text, size=size)
    return p


def add_bullet(doc, text, *, after=3):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    add_run(p, text, size=10)
    return p


def add_heading(doc, text, level=1, before=10, after=5):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    add_run(p, text, bold=True, color=BLACK, size=14 if level == 1 else 11.5)
    return p


def add_label_value_row(table, label, value):
    row = table.add_row()
    set_cell_text(row.cells[0], label, bold=True, color=RGBColor(74, 86, 102), size=9)
    set_cell_text(row.cells[1], value, color=BODY, size=9.5)
    return row


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(paragraph, "Autozentic  |  Confidential  |  ", color=RGBColor(100, 110, 123), size=8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY

    for style_name, size in (("Title", 28), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    add_run(hp, "AUTOZENTIC", bold=True, color=BLACK, size=11)
    add_run(hp, "     DIGITAL AUTOMATION SOLUTIONS", bold=True, color=RGBColor(93, 105, 119), size=8)
    hp_pr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), TEAL)
    pbdr.append(bottom)
    hp_pr.append(pbdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(4)
    add_run(title, "AI Calling Agent Proposal", bold=True, color=BLACK, size=27)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    add_run(subtitle, "AI-led customer conversations and appointment booking", color=RGBColor(83, 96, 112), size=11.5)

    meta = doc.add_table(rows=0, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    meta.columns[0].width = Inches(1.35)
    meta.columns[1].width = Inches(5.15)
    set_table_borders(meta, color="FFFFFF", size="0")
    add_label_value_row(meta, "Prepared by", "Autozentic")
    add_label_value_row(meta, "Location", "Jagatpura, Jaipur")
    add_label_value_row(meta, "Prepared for", "Sekol Tiles, Gujarat")
    add_label_value_row(meta, "Date", "8 September 2026")
    for row in meta.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=45, bottom=45, start=0, end=120)

    add_heading(doc, "Proposal Summary", before=22, after=6)
    add_body(doc, "Autozentic will design, configure and hand over an AI calling agent with appointment booking support. The solution is intended to help the client respond faster, keep customer context across the conversation and make follow-up calls more consistent.", after=7)

    add_heading(doc, "Included Solution", before=16, after=7)
    scope = doc.add_table(rows=1, cols=2)
    scope.alignment = WD_TABLE_ALIGNMENT.LEFT
    scope.autofit = False
    scope.columns[0].width = Inches(1.72)
    scope.columns[1].width = Inches(4.78)
    set_table_borders(scope)
    set_cell_shading(scope.rows[0].cells[0], NAVY)
    set_cell_shading(scope.rows[0].cells[1], NAVY)
    set_cell_text(scope.rows[0].cells[0], "Solution", bold=True, color=WHITE, size=9.2)
    set_cell_text(scope.rows[0].cells[1], "Included capability", bold=True, color=WHITE, size=9.2)
    set_repeat_table_header(scope.rows[0])
    rows = [
        ("AI calling agent", "AI-led inbound and outbound conversations, customer context, purpose-led replies, call outcomes, transcripts and CRM call logs."),
        ("Appointment booking", "Date and time collection, availability checks, confirmed booking creation and appointment visibility in the CRM."),
        ("Deployment and support", "One-year VPS plan, one-year bug support and handover guidance for the delivered solution."),
    ]
    for index, (label, description) in enumerate(rows):
        row = scope.add_row()
        if index % 2 == 1:
            set_cell_shading(row.cells[0], LIGHT_GRAY)
            set_cell_shading(row.cells[1], LIGHT_GRAY)
        set_cell_text(row.cells[0], label, bold=True, color=BLACK, size=9.2)
        set_cell_text(row.cells[1], description, size=9.2)

    add_heading(doc, "Delivery Approach", before=16, after=5)
    add_bullet(doc, "Confirm business workflows, channel access and appointment rules.")
    add_bullet(doc, "Configure the calling agent, CRM logging and appointment booking flow.")
    add_bullet(doc, "Run integration testing, conversation testing and client review.")
    add_bullet(doc, "Complete handover and begin the one-year bug-support period after delivery.")

    doc.add_page_break()
    add_heading(doc, "Commercials", before=0, after=7)
    add_body(doc, "The complete AI calling agent solution is offered at a professional fee of INR 16,000. Goods and Services Tax is charged at 18% as applicable.", after=10)

    pricing = doc.add_table(rows=1, cols=3)
    pricing.alignment = WD_TABLE_ALIGNMENT.LEFT
    pricing.autofit = False
    pricing.columns[0].width = Inches(3.6)
    pricing.columns[1].width = Inches(1.45)
    pricing.columns[2].width = Inches(1.45)
    set_table_borders(pricing)
    headers = ("Description", "Amount", "GST 18%")
    for cell, text in zip(pricing.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        set_cell_text(cell, text, bold=True, color=WHITE, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER if text != "Description" else WD_ALIGN_PARAGRAPH.LEFT)
    set_repeat_table_header(pricing.rows[0])
    data = [
        ("Complete solution", "INR 16,000", "INR 2,880"),
        ("Total payable", "INR 18,880", "Included"),
    ]
    for index, values in enumerate(data):
        row = pricing.add_row()
        if index == 1:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_TEAL)
        for col, value in enumerate(values):
            set_cell_text(row.cells[col], value, bold=index == 1, color=BLACK if index == 1 else BODY, size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT if col > 0 else WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "Payment Schedule", before=16, after=6)
    payment = doc.add_table(rows=1, cols=3)
    payment.alignment = WD_TABLE_ALIGNMENT.LEFT
    payment.autofit = False
    payment.columns[0].width = Inches(1.7)
    payment.columns[1].width = Inches(3.1)
    payment.columns[2].width = Inches(1.7)
    set_table_borders(payment)
    for cell, text in zip(payment.rows[0].cells, ("Milestone", "Payment condition", "Amount incl. GST")):
        set_cell_shading(cell, NAVY)
        set_cell_text(cell, text, bold=True, color=WHITE, size=9.1)
    set_repeat_table_header(payment.rows[0])
    payment_rows = [
        ("Advance", "50% before project commencement", "INR 9,440"),
        ("Completion", "50% after completion and handover", "INR 9,440"),
    ]
    for index, values in enumerate(payment_rows):
        row = payment.add_row()
        if index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
        for col, value in enumerate(values):
            set_cell_text(row.cells[col], value, bold=col == 0, color=BLACK if col == 0 else BODY, size=9.2, align=WD_ALIGN_PARAGRAPH.RIGHT if col == 2 else WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "Client Paid Usage Charges", before=16, after=6)
    add_body(doc, "AI calling usage charges are paid directly by the client to the calling provider and are separate from Autozentic's professional fee:", after=5)
    add_bullet(doc, "AI calling and telephony charges billed per minute by the calling provider.")

    doc.add_page_break()
    add_heading(doc, "Support and Responsibilities", before=0, after=7)
    add_heading(doc, "Included support", level=2, before=4, after=5)
    add_body(doc, "Autozentic will provide one year of bug support for the delivered features and one year of VPS hosting as part of this proposal. Bug support covers defects in the delivered implementation. New features, new campaign content, provider charges and workflow changes outside the agreed scope may be quoted separately.", after=9)

    add_heading(doc, "Client responsibilities", level=2, before=12, after=5)
    add_bullet(doc, "Provide timely access to the telephony, CRM and other required accounts.")
    add_bullet(doc, "Complete provider verification, business approvals and payment of all third-party usage charges.")
    add_bullet(doc, "Provide approved brand content, FAQs, campaign copy, appointment rules and escalation contacts.")
    add_bullet(doc, "Review and approve the configured workflows during testing and handover.")

    add_heading(doc, "Acceptance and handover", level=2, before=12, after=5)
    add_body(doc, "Completion will be treated as the point at which the agreed features are configured, tested and demonstrated to the client. The remaining 50% payment becomes due after completion and handover. The one-year bug-support period begins on the handover date.", after=9)

    add_heading(doc, "Approval", before=16, after=6)
    add_body(doc, "To begin, the client may confirm acceptance of this proposal and pay the 50% advance shown above. The project start date will be scheduled after the advance and required account access are received.", after=16)

    sign = doc.add_table(rows=1, cols=2)
    sign.alignment = WD_TABLE_ALIGNMENT.LEFT
    sign.autofit = False
    sign.columns[0].width = Inches(3.1)
    sign.columns[1].width = Inches(3.1)
    set_table_borders(sign, color="FFFFFF", size="0")
    set_cell_text(sign.rows[0].cells[0], "For Autozentic\n\nName: ____________________\nSignature: _______________", size=9.5)
    set_cell_text(sign.rows[0].cells[1], "For Client\n\nName: ____________________\nSignature: _______________", size=9.5)
    for cell in sign.rows[0].cells:
        set_cell_margins(cell, top=0, bottom=0, start=0, end=180)

    doc.core_properties.title = "AI Calling Agent Proposal - Autozentic"
    doc.core_properties.subject = "Proposal for AI calling and appointment booking"
    doc.core_properties.author = "Autozentic"
    doc.core_properties.keywords = "Autozentic, AI calling, appointment booking, proposal"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
